"""
File Processor - Document Extraction Module
Handles PDF, DOCX, TXT, and Image file processing
"""

import os
import io
from typing import Optional, Union
from dataclasses import dataclass


@dataclass
class ProcessedFile:
    """Processed file result"""
    filename: str
    content: str
    file_type: str
    page_count: Optional[int] = None


class FileProcessor:
    """Processor for extracting text from various file formats"""
    
    SUPPORTED_TYPES = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".txt": "text",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image",
        ".gif": "image"
    }
    
    def __init__(self):
        self.processors = {
            "pdf": self._process_pdf,
            "docx": self._process_docx,
            "text": self._process_text,
            "image": self._process_image
        }
    
    def process_file(self, uploaded_file) -> str:
        """
        Process an uploaded file and extract text
        
        Args:
            uploaded_file: Streamlit uploaded file object
            
        Returns:
            Extracted text content
        """
        # Get file extension
        filename = uploaded_file.name
        ext = os.path.splitext(filename)[1].lower()
        
        # Determine file type
        file_type = self.SUPPORTED_TYPES.get(ext, "unknown")
        
        if file_type == "unknown":
            return f"[Unsupported file type: {ext}]"
        
        # Get processor
        processor = self.processors.get(file_type)
        if not processor:
            return "[No processor available]"
        
        try:
            return processor(uploaded_file)
        except Exception as e:
            return f"[Error processing {filename}: {str(e)}]"
    
    def _process_pdf(self, uploaded_file) -> str:
        """Extract text from PDF"""
        try:
            import PyPDF2
            
            # Read file content
            uploaded_file.seek(0)
            pdf_bytes = uploaded_file.read()
            
            # Create PDF reader
            from io import BytesIO
            pdf_file = BytesIO(pdf_bytes)
            reader = PyPDF2.PdfReader(pdf_file)
            
            # Extract text from all pages
            text_parts = []
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(f"[Page {page_num + 1}]\n{text}")
            
            return "\n\n".join(text_parts)
            
        except ImportError:
            return "[PyPDF2 not installed - cannot process PDF]"
        except Exception as e:
            return f"[PDF processing error: {str(e)}]"
    
    def _process_docx(self, uploaded_file) -> str:
        """Extract text from DOCX"""
        try:
            import docx
            
            # Read file content
            uploaded_file.seek(0)
            doc_bytes = uploaded_file.read()
            
            # Create document from bytes
            from io import BytesIO
            doc_file = BytesIO(doc_bytes)
            doc = docx.Document(doc_file)
            
            # Extract text from all paragraphs
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)
            
            return "\n".join(text_parts)
            
        except ImportError:
            return "[python-docx not installed - cannot process DOCX]"
        except Exception as e:
            return f"[DOCX processing error: {str(e)}]"
    
    def _process_text(self, uploaded_file) -> str:
        """Process plain text file"""
        uploaded_file.seek(0)
        return uploaded_file.read().decode('utf-8', errors='ignore')
    
    def _process_image(self, uploaded_file) -> str:
        """Process image file - extract text via OCR"""
        try:
            from PIL import Image
            import pytesseract
            
            # Read image
            uploaded_file.seek(0)
            image = Image.open(uploaded_file)
            
            # Perform OCR
            text = pytesseract.image_to_string(image)
            
            return f"[OCR Extracted from image]\n{text}"
            
        except ImportError:
            return "[Image OCR not available - install pytesseract]"
        except Exception as e:
            return f"[Image processing error: {str(e)}]"
    
    def process_multiple(self, uploaded_files) -> str:
        """
        Process multiple files and combine content
        
        Args:
            uploaded_files: List of uploaded file objects
            
        Returns:
            Combined text from all files
        """
        contents = []
        
        for file in uploaded_files:
            content = self.process_file(file)
            contents.append(f"=== {file.name} ===\n{content}")
        
        return "\n\n".join(contents)
    
    def get_file_info(self, uploaded_file) -> dict:
        """
        Get basic file information
        
        Args:
            uploaded_file: Uploaded file object
            
        Returns:
            Dictionary with file info
        """
        return {
            "name": uploaded_file.name,
            "size": uploaded_file.size,
            "type": uploaded_file.type,
            "extension": os.path.splitext(uploaded_file.name)[1].lower()
        }


# Standalone function
def extract_text(uploaded_file) -> str:
    """
    Extract text from uploaded file
    
    Args:
        uploaded_file: Streamlit uploaded file
        
    Returns:
        Extracted text
    """
    processor = FileProcessor()
    return processor.process_file(uploaded_file)


if __name__ == "__main__":
    # Test
    print("FileProcessor module loaded")