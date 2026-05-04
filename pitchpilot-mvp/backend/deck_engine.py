"""
Deck Engine - PowerPoint Generation Module
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


class DeckEngine:
    """Engine for generating PowerPoint presentations"""
    
    TITLE_SLIDE_LAYOUT = 0
    TITLE_CONTENT_LAYOUT = 1
    
    PRIMARY_COLOR = RGBColor(0, 51, 102)
    ACCENT_COLOR = RGBColor(0, 120, 212)
    TEXT_COLOR = RGBColor(51, 51, 51)
    
    def __init__(self):
        self.slide_templates = {
            1: {"title": "The Hook"},
            2: {"title": "Empathy"},
            3: {"title": "Opportunity"},
            4: {"title": "Solution"},
            5: {"title": "Market"},
            6: {"title": "Business Model"},
            7: {"title": "Competition"},
            8: {"title": "Technology"},
            9: {"title": "Traction"},
            10: {"title": "The Ask"}
        }
    
    def generate_presentation(self, deck_data: dict, output_path: str = None) -> str:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # Title slide
        self._create_title_slide(prs, deck_data)
        
        # Content slides
        for slide_num in range(1, 11):
            slide_key = f"slide_{slide_num}"
            content = deck_data.get(slide_key, "")
            self._create_content_slide(prs, slide_num, content)
        
        if not output_path:
            company_name = deck_data.get("company_name", "pitch_deck")
            company_name = "".join(c for c in company_name if c.isalnum() or c in " -_").strip()
            output_dir = os.path.dirname(os.path.abspath(__file__)) if os.path.dirname(__file__) else "."
            output_path = os.path.join(output_dir, f"{company_name}_pitch_deck.pptx")
        
        prs.save(output_path)
        return output_path
    
    def _create_title_slide(self, prs: Presentation, deck_data: dict):
        slide_layout = prs.slide_layouts[self.TITLE_SLIDE_LAYOUT]
        slide = prs.slides.add_slide(slide_layout)
        
        company_name = deck_data.get("company_name", "")
        if company_name:
            title = slide.shapes.title
            title.text = company_name
        
        subtitle = slide.placeholders[1]
        founder_name = deck_data.get("founder_name", "")
        subtitle_text = founder_name + "\n" + deck_data.get("slide_1_hook", "") if founder_name else deck_data.get("slide_1_hook", "")
        subtitle.text = subtitle_text
    
    def _create_content_slide(self, prs: Presentation, slide_num: int, content: str):
        slide_layout = prs.slide_layouts[self.TITLE_CONTENT_LAYOUT]
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        title.text = self.slide_templates.get(slide_num, {"title": f"Slide {slide_num}"})["title"]
        
        for paragraph in title.text_frame.paragraphs:
            paragraph.font.size = Pt(32)
            paragraph.font.bold = True
            paragraph.font.color.rgb = self.PRIMARY_COLOR
        
        content_placeholder = slide.placeholders[1]
        if content:
            content_placeholder.text_frame.clear()
            p = content_placeholder.text_frame.paragraphs[0]
            p.text = content
            p.font.size = Pt(18)
            p.font.color.rgb = self.TEXT_COLOR